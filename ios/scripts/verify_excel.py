#!/usr/bin/env python3
"""Read XCTest's actual Swift-produced workbook with an independent OOXML implementation."""
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile
import openpyxl

fixtures = list((Path.home() / 'Library/Developer/CoreSimulator/Devices').glob('*/data/Containers/Data/Application/*/tmp/IZZ-ExcelRegression.xlsx'))
if not fixtures:
    raise SystemExit('XCTest did not produce the Excel regression fixture')
file = max(fixtures, key=lambda p: p.stat().st_mtime)
with ZipFile(file) as archive:
    assert archive.testzip() is None, 'ZIP CRC failure'
wb = openpyxl.load_workbook(file)
ws = wb['Delivery Schedule']
assert ws['A1'].value == 'IZZ Delivery - Complete Delivery Schedule'
assert ws['B5'].value == 'My Home — Home & Road'
assert ws['B6'].value == ws['B5'].value
assert [ws.cell(r, 2).value for r in range(9, 14)] == [50, 1, 1, 1, 47]
assert ws['B14'].value == .02
assert ws['B15'].value == 63.75
assert ws['B16'].value == 99.1 and ws['B17'].value == 5.2
assert ws['B16'].number_format == '"SGD "#,##0.00'
headers = ['Stop','Postal Code','Block','Area','Arrival','Distance From Previous','Base Drive','Traffic Buffer','Planned Travel','Delivery','Leave','Cumulative KM']
assert [ws.cell(20, c).value for c in range(1,13)] == headers
assert ws['A21'].value == 'START' and ws['A72'].value == 'END'
assert ws['B21'].value == '001234' and ws['B72'].value == '001234'
assert ws['C21'].value == 'My Home' and ws['C72'].value == 'My Home'
assert ws['B22'].value == '012345' and ws['B22'].data_type == 's'
assert ws['C22'].value == '=1+1' and ws['C22'].data_type == 's'
assert ws['D22'].value == '道路 & Road'
assert [ws.cell(r,1).value for r in range(22,72)] == list(range(1,51))
assert ws['E21'].value == datetime(2026,10,4,23,59)
assert ws['E22'].value.date() > ws['E21'].value.date()
assert ws['E22'].number_format == 'yyyy-mm-dd hh:mm'
assert ws['F22'].number_format == '0.00 "KM"'
assert ws['G22'].number_format == '0.0 "min"'
assert ws.freeze_panes == 'A21' and ws['A20'].font.bold
assert ws['A20'].fill.fgColor.rgb == 'FF081326'
assert ws.column_dimensions['D'].width >= 28
print('Verified Swift XLSX: 50 stops, saved summary/Home, midnight, leading zeroes, literal text, styles and ZIP CRC.')
